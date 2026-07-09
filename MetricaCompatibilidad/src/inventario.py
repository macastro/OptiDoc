"""
Examina un documento antes de convertirlo para saber con qué complejidad
nos vamos a encontrar: tablas, estilos, fuentes, encabezados/pies, imágenes,
listas, columnas y cambios controlados. Con eso calculamos un "esfuerzo previsto"
que sirve de predicción antes de arrancar la conversión.

Para las fuentes distinguimos entre las que están DECLARADAS en el XML interno
(que incluye fuentes que Word agrega por defecto) y las que el autor realmente
USÓ en el texto — esta última señal es más limpia. También sabemos qué fuentes
LibreOffice reemplaza de forma transparente sin alterar la maquetación, así que
a esas no las penalizamos.
"""
from __future__ import annotations

import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

# Namespaces XML que parsea este módulo.
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
STYLE_ODF = "urn:oasis:names:tc:opendocument:xmlns:style:1.0"

# Fuentes habituales que suelen estar disponibles en los dos motores.
FUENTES_SEGURAS = {
    "Arial", "Calibri", "Times New Roman", "Liberation Serif", "Liberation Sans",
    "Liberation Mono", "DejaVu Sans", "Carlito", "Caladea", "Cambria", "Verdana",
}

# LibreOffice reemplaza estas fuentes por equivalentes de mismas métricas —
# el texto fluye igual, así que no las penalizamos en el puntaje.
SUSTITUCIONES_SEGURAS = {
    "Arial": "Liberation Sans",
    "Helvetica": "Liberation Sans",
    "Times New Roman": "Liberation Serif",
    "Courier New": "Liberation Mono",
    "Calibri": "Carlito",
    "Cambria": "Caladea",
}

# Fuentes que Word y LibreOffice añaden solos aunque el autor no las pidió —
# las ignoramos para no inflar artificialmente la complejidad.
FUENTES_DEFAULT = {
    "Courier", "Symbol", "ＭＳ ゴシック", "ＭＳ 明朝", "MS Gothic", "MS Mincho",
}


def _q(tag: str) -> str:
    return f"{{{W}}}{tag}"


def _root_document(ruta: str | Path):
    with zipfile.ZipFile(ruta) as z:
        if "word/document.xml" in z.namelist():
            return ET.fromstring(z.read("word/document.xml"))
    return None


def _contar_local(root, local: str) -> int:
    """Cuenta elementos XML por nombre de etiqueta sin importar el namespace."""
    if root is None:
        return 0
    return sum(1 for el in root.iter() if el.tag.split("}")[-1] == local)


def fuentes_docx(ruta: str | Path) -> set[str]:
    """Lee las fuentes declaradas en fontTable.xml (puede incluir fuentes que Word agregó solo)."""
    fuentes: set[str] = set()
    with zipfile.ZipFile(ruta) as z:
        if "word/fontTable.xml" in z.namelist():
            root = ET.fromstring(z.read("word/fontTable.xml"))
            for f in root.findall(_q("font")):
                nombre = f.get(_q("name"))
                if nombre:
                    fuentes.add(nombre)
    return fuentes


def fuentes_usadas_docx(doc: Document) -> set[str]:
    """Recoge las fuentes que el autor asignó explícitamente en el texto, tablas, encabezados y pies."""
    usadas: set[str] = set()

    def _de_parrafos(parrafos):
        for p in parrafos:
            for r in p.runs:
                if r.font is not None and r.font.name:
                    usadas.add(r.font.name)

    _de_parrafos(doc.paragraphs)
    for t in doc.tables:
        for fila in t.rows:
            for celda in fila.cells:
                _de_parrafos(celda.paragraphs)
    for s in doc.sections:
        _de_parrafos(s.header.paragraphs)
        _de_parrafos(s.footer.paragraphs)
    return usadas


def fuentes_odt(ruta: str | Path) -> set[str]:
    """Extrae las fuentes declaradas en el archivo ODF de destino."""
    fuentes: set[str] = set()
    with zipfile.ZipFile(ruta) as z:
        for miembro in ("styles.xml", "content.xml"):
            if miembro in z.namelist():
                root = ET.fromstring(z.read(miembro))
                for ff in root.iter(f"{{{STYLE_ODF}}}font-face"):
                    nombre = ff.get(f"{{{STYLE_ODF}}}name")
                    if nombre:
                        fuentes.add(nombre)
    return fuentes


def inventariar(ruta_docx: str | Path) -> dict:
    """Examina el documento y estima su complejidad de conversión."""
    ruta_docx = Path(ruta_docx)
    doc = Document(str(ruta_docx))
    root = _root_document(ruta_docx)

    n_tablas = len(doc.tables)
    n_parrafos = len(doc.paragraphs)
    estilos = {p.style.name for p in doc.paragraphs if p.style is not None}

    # Las imágenes en OOXML se representan como 'blip' dentro de DrawingML.
    n_imagenes = _contar_local(root, "blip")
    # Detectamos listas tanto por su estilo ('List Bullet'/'List Number') como por la propiedad numPr.
    listas_por_estilo = sum(
        1 for p in doc.paragraphs
        if p.style is not None and p.style.name and p.style.name.startswith("List")
    )
    n_listas = max(listas_por_estilo, _contar_local(root, "numPr"))
    control_cambios = (_contar_local(root, "ins") + _contar_local(root, "del")) > 0
    n_hipervinculos = _contar_local(root, "hyperlink")
    n_objetos = _contar_local(root, "object") + _contar_local(root, "OLEObject")

    n_secciones = len(doc.sections)
    n_encabezados = sum(
        1 for s in doc.sections
        if not s.header.is_linked_to_previous
        and any(p.text.strip() for p in s.header.paragraphs)
    )
    n_pies = sum(
        1 for s in doc.sections
        if not s.footer.is_linked_to_previous
        and any(p.text.strip() for p in s.footer.paragraphs)
    )
    columnas_max = 1
    for s in doc.sections:
        cols = s._sectPr.find(_q("cols"))
        if cols is not None:
            num = cols.get(_q("num"))
            if num and num.isdigit():
                columnas_max = max(columnas_max, int(num))

    fuentes_declaradas = fuentes_docx(ruta_docx)
    fuentes_usadas = fuentes_usadas_docx(doc)
    # Solo nos interesan las fuentes que el autor eligió, no las que Word mete por su cuenta.
    fuentes_interes = (fuentes_usadas or fuentes_declaradas) - FUENTES_DEFAULT
    fuentes_no_seguras = sorted(fuentes_interes - FUENTES_SEGURAS - set(SUSTITUCIONES_SEGURAS))

    inv = {
        "archivo": ruta_docx.name,
        "n_parrafos": n_parrafos,
        "n_tablas": n_tablas,
        "n_estilos_distintos": len(estilos),
        "n_fuentes_declaradas": len(fuentes_declaradas),
        "fuentes_declaradas": sorted(fuentes_declaradas),
        "fuentes_usadas": sorted(fuentes_usadas),
        "fuentes_no_seguras": fuentes_no_seguras,
        "n_imagenes": n_imagenes,
        "n_objetos_incrustados": n_objetos,
        "n_listas": n_listas,
        "n_hipervinculos": n_hipervinculos,
        "n_secciones": n_secciones,
        "columnas_max": columnas_max,
        "n_encabezados": n_encabezados,
        "n_pies_pagina": n_pies,
        "control_cambios": control_cambios,
    }

    # Puntuación de riesgo: cada elemento que suele dar problemas al convertir suma puntos.
    riesgo = (
        2.0 * n_tablas
        + 4.0 * n_objetos
        + 2.0 * n_imagenes
        + 1.0 * n_listas
        + 3.0 * max(columnas_max - 1, 0)
        + 2.0 * len(fuentes_no_seguras)
        + (4.0 if control_cambios else 0.0)
    )
    esfuerzo_previsto = round(min(100.0, riesgo * 4.0), 1)
    if esfuerzo_previsto < 25:
        complejidad = "baja"
    elif esfuerzo_previsto < 60:
        complejidad = "media"
    else:
        complejidad = "alta"

    inv["esfuerzo_previsto"] = esfuerzo_previsto
    inv["complejidad"] = complejidad
    return inv


def fuentes_para_puntaje(ruta_docx: str | Path) -> set[str]:
    """Devuelve las fuentes relevantes del documento, filtrando las que Word inserta automáticamente."""
    return fuentes_docx(ruta_docx) - FUENTES_DEFAULT
