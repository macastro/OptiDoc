"""
generar_sintetico.py — Genera un corpus sintético de prueba (python-docx).

Crea dos documentos que reproducen características críticas de los documentos
institucionales (Sección 6, Fase 1), para ejercitar el pipeline sin depender aún
de documentos reales:

  01_memo_simple.docx     -> texto y tabla simples, fuentes comunes (complejidad baja)
  02_informe_complejo.docx-> multipágina con tablas (celdas combinadas), listas,
                             imagen incrustada, líneas tipo formulario con
                             tabuladores, sección a dos columnas y una fuente poco
                             común (complejidad alta)

Uso:  python corpus/generar_sintetico.py
Salida: corpus/sinteticos/*.docx
"""
from __future__ import annotations

from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SALIDA = Path(__file__).resolve().parent / "sinteticos"

_PARRAFO = (
    "La interoperabilidad entre suites ofimáticas es un requisito creciente en "
    "entornos institucionales donde conviven herramientas propietarias y de "
    "software libre. La fidelidad de formato al intercambiar documentos determina "
    "el esfuerzo de reformateo y, por tanto, el costo real de la fricción. Este "
    "documento sintético reproduce elementos típicos cuyo mapeo entre OOXML y ODF "
    "conviene evaluar de forma objetiva y reproducible."
)


def _set_columnas(section, num: int = 2, space_twips: int = 425) -> None:
    """Configura el número de columnas de una sección (XML directo)."""
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space_twips))


def _campo_pagina(paragraph) -> None:
    """Inserta un campo PAGE (número de página) en un párrafo."""
    run = paragraph.add_run()
    ini = OxmlElement("w:fldChar"); ini.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fin = OxmlElement("w:fldChar"); fin.set(qn("w:fldCharType"), "end")
    run._r.append(ini); run._r.append(instr); run._r.append(fin)


def _grafico(ruta: Path) -> Path:
    """Genera una imagen (gráfico de barras) para incrustar."""
    fig, ax = plt.subplots(figsize=(5, 3))
    ax.bar(["Word", "LibreOffice", "Google Docs"], [42, 31, 27],
           color=["#2E75B6", "#43A047", "#F9A825"])
    ax.set_ylabel("Uso (%)")
    ax.set_title("Uso declarado de suites (ejemplo)")
    fig.tight_layout()
    ruta.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(ruta, dpi=120)
    plt.close(fig)
    return ruta


def _encabezado_pie(doc, titulo: str) -> None:
    sec = doc.sections[0]
    enc = sec.header.paragraphs[0]
    enc.text = titulo
    pie = sec.footer.paragraphs[0]
    pie.text = "Documento sintético · Página "
    _campo_pagina(pie)


def memo_simple(ruta: Path) -> None:
    doc = Document()
    _encabezado_pie(doc, "ESPOL · Memorando (sintético)")
    doc.add_heading("Memorando", level=0)
    p = doc.add_paragraph()
    p.add_run("Para: ").bold = True
    p.add_run("Coordinación Académica\n")
    p.add_run("De: ").bold = True
    p.add_run("Secretaría\n")
    p.add_run("Asunto: ").bold = True
    p.add_run("Prueba de compatibilidad de formato")
    doc.add_paragraph(_PARRAFO)
    doc.add_heading("Resumen", level=1)
    tabla = doc.add_table(rows=3, cols=2)
    tabla.style = "Table Grid"
    datos = [("Concepto", "Valor"), ("Documentos", "3"), ("Estado", "En revisión")]
    for i, (a, b) in enumerate(datos):
        tabla.rows[i].cells[0].text = a
        tabla.rows[i].cells[1].text = b
    doc.add_paragraph("Atentamente,")
    doc.save(ruta)


def informe_complejo(ruta: Path) -> None:
    doc = Document()
    _encabezado_pie(doc, "ESPOL · Informe técnico (sintético)")
    doc.add_heading("Informe técnico de compatibilidad", level=0)

    # Línea tipo formulario con tabuladores y línea de puntos
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.add_run("Responsable: ").bold = True
    p.add_run("\tFecha: __________")

    doc.add_heading("1. Introducción", level=1)
    for _ in range(3):
        doc.add_paragraph(_PARRAFO)

    # Fuente poco común para ejercitar la dimensión de fuentes
    p = doc.add_paragraph()
    r = p.add_run("Texto en una fuente poco común (Garamond) para evaluar sustitución.")
    r.font.name = "EB Garamond"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    doc.add_heading("2. Lista de verificación", level=1)
    for item in ["Tablas con celdas combinadas", "Encabezados y pies",
                 "Imágenes incrustadas", "Numeración y viñetas", "Columnas múltiples"]:
        doc.add_paragraph(item, style="List Bullet")
    for paso in ["Inventariar características", "Convertir a ODF", "Comparar fidelidad"]:
        doc.add_paragraph(paso, style="List Number")

    doc.add_heading("3. Tabla con celdas combinadas", level=1)
    tabla = doc.add_table(rows=4, cols=3)
    tabla.style = "Table Grid"
    # Encabezado combinado en la primera fila
    a = tabla.cell(0, 0); b = tabla.cell(0, 2)
    a.merge(b).text = "Resultados de la conversión (encabezado combinado)"
    enc = ["Elemento", "Origen", "Destino"]
    for j, t in enumerate(enc):
        tabla.cell(1, j).text = t
    filas = [("Tablas", "2", "2"), ("Fuentes", "3", "2")]
    for i, fila in enumerate(filas, start=2):
        for j, t in enumerate(fila):
            tabla.cell(i, j).text = t

    doc.add_heading("4. Figura incrustada", level=1)
    img = _grafico(SALIDA / "_grafico.png")
    doc.add_picture(str(img), width=Inches(4.2))

    # Relleno para forzar varias páginas (cambios de paginación detectables)
    doc.add_heading("5. Desarrollo", level=1)
    for _ in range(8):
        doc.add_paragraph(_PARRAFO)

    # Sección final a dos columnas
    doc.add_section()
    _set_columnas(doc.sections[-1], num=2)
    doc.add_heading("6. Anexo a dos columnas", level=1)
    for _ in range(4):
        doc.add_paragraph(_PARRAFO)

    doc.save(ruta)


def main() -> None:
    SALIDA.mkdir(parents=True, exist_ok=True)
    memo_simple(SALIDA / "01_memo_simple.docx")
    informe_complejo(SALIDA / "02_informe_complejo.docx")
    # Limpia la imagen temporal
    tmp = SALIDA / "_grafico.png"
    if tmp.exists():
        tmp.unlink()
    print("Corpus sintético generado en:", SALIDA)
    for f in sorted(SALIDA.glob("*.docx")):
        print("  -", f.name)


if __name__ == "__main__":
    main()
